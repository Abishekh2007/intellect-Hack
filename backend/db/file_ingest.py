"""File ingestion.

Converts uploaded data files (CSV, XLSX, PDF, DOCX, JSON) into queryable
SQLite tables. Used to power multi-turn questions over a user's own data.
"""

from __future__ import annotations

import csv
import io
import json
import re
import sqlite3
from pathlib import Path
from typing import Any

from config import get_settings

ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xls", ".pdf", ".docx", ".json"}

_TABLE_NAME_RE = re.compile(r"[^A-Za-z0-9_]")


def sanitize_table_name(name: str) -> str:
    cleaned = _TABLE_NAME_RE.sub("_", name)
    if not cleaned:
        cleaned = "uploaded_data"
    if cleaned[0].isdigit():
        cleaned = "t_" + cleaned
    return cleaned[:60]


def _detect_type(value: str) -> str:
    v = value.strip()
    if not v:
        return "TEXT"
    if re.fullmatch(r"-?\d+", v):
        return "INTEGER"
    if re.fullmatch(r"-?\d+\.\d+", v):
        return "REAL"
    return "TEXT"


def _csv_rows(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return [row for row in reader if any(c.strip() for c in row)]


def _json_rows(content: bytes) -> list[list[str]]:
    data = json.loads(content.decode("utf-8", errors="replace"))
    records: list[dict[str, Any]] = []
    if isinstance(data, dict):
        for v in data.values():
            if isinstance(v, list) and v and isinstance(v[0], dict):
                records.extend(v)
            elif isinstance(v, list):
                records.extend({"value": item} for item in v)
        if not records and all(not isinstance(v, (dict, list)) for v in data.values()):
            records = [data]
    elif isinstance(data, list):
        records = data

    # Flatten nested dicts one level.
    flat: list[dict[str, Any]] = []
    for rec in records:
        if not isinstance(rec, dict):
            flat.append({"value": rec})
            continue
        out: dict[str, Any] = {}
        for k, v in rec.items():
            if isinstance(v, dict):
                for k2, v2 in v.items():
                    out[f"{k}_{k2}"] = v2
            elif isinstance(v, (list, tuple)):
                out[k] = ", ".join(str(x) for x in v)
            else:
                out[k] = v
        flat.append(out)

    if not flat:
        return []
    columns = list(flat[0].keys())
    return [[str(rec.get(c, "")) for c in columns] for rec in flat]


def _excel_rows(content: bytes) -> list[list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = wb.active
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(["" if v is None else str(v) for v in row])
    return rows


def _pdf_rows(content: bytes) -> list[list[str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    lines: list[list[str]] = []
    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for line in text.splitlines():
            lines.append([str(page_num), line])
    return lines


def _docx_rows(content: bytes) -> list[list[str]]:
    import docx

    doc = docx.Document(io.BytesIO(content))
    rows: list[list[str]] = []
    for para in doc.paragraphs:
        if para.text.strip():
            rows.append([para.text])
    for table in doc.tables:
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
    return rows


def extract_rows(filename: str, content: bytes) -> tuple[list[str], list[list[str]]]:
    """Return (columns, rows) from an uploaded file."""
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        rows = _csv_rows(content)
        if not rows:
            return [], []
        columns = rows[0]
        data = rows[1:]
    elif ext == ".json":
        # Straight to the columns-aware reader. The previous version called
        # _json_rows, computed a `columns` value, discarded both, and returned
        # _json_rows_with_columns anyway — parsing the same payload three
        # times to reach the answer the third call already had.
        return _json_rows_with_columns(content)
    elif ext in (".xlsx", ".xls"):
        rows = _excel_rows(content)
        if not rows:
            return [], []
        columns = rows[0]
        data = rows[1:]
    elif ext == ".pdf":
        rows = _pdf_rows(content)
        return ["page_number", "line"], rows
    elif ext == ".docx":
        rows = _docx_rows(content)
        return ["line"], rows
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    return columns, data


def _json_rows_with_columns(content: bytes) -> tuple[list[str], list[list[str]]]:
    data = json.loads(content.decode("utf-8", errors="replace"))
    records: list[dict[str, Any]] = []
    if isinstance(data, dict):
        for v in data.values():
            if isinstance(v, list) and v and isinstance(v[0], dict):
                records.extend(v)
    elif isinstance(data, list):
        records = data
    if not records or not isinstance(records[0], dict):
        return ["value"], [[str(r)] for r in records] if records else []
    columns = list(records[0].keys())
    rows = []
    for rec in records:
        if isinstance(rec, dict):
            rows.append([rec.get(c, "") for c in columns])
        else:
            rows.append([rec])
    return columns, rows


def _infer_types(columns: list[str], data: list[list[str]]) -> list[str]:
    """Infer a SQLite column type per column from the sampled values.

    Every column used to come back TEXT: the accumulator started at "TEXT",
    and the first branch that saw a number hit ``if types[i] == "TEXT":
    continue`` and left it there forever. Numeric uploads therefore sorted and
    summed as strings — "9" ranked above "100" in every ORDER BY.

    Blank cells are ignored rather than forcing TEXT, so one empty cell in a
    numeric column no longer downgrades the whole column.
    """
    types: list[str] = []
    for col_idx in range(len(columns)):
        seen: set[str] = set()
        for row in data:
            if col_idx >= len(row):
                continue
            value = row[col_idx]
            if value is None or not str(value).strip():
                continue  # blanks are unknown, not text
            seen.add(_detect_type(str(value)))
            if "TEXT" in seen:
                break  # one real string settles it
        if not seen or "TEXT" in seen:
            types.append("TEXT")
        elif seen == {"INTEGER"}:
            types.append("INTEGER")
        else:
            types.append("REAL")  # INTEGER mixed with REAL widens to REAL
    return types


def ingest_upload(filename: str, content: bytes, conn: sqlite3.Connection) -> dict[str, Any]:
    """Create a table from an uploaded file and return table info."""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}")
    if len(content) > 25 * 1024 * 1024:
        raise ValueError("File too large (max 25 MB).")

    columns, data = extract_rows(filename, content)
    if not columns:
        raise ValueError("File appears to be empty.")
    if len(data) > 100_000:
        data = data[:100_000]

    table_name = sanitize_table_name(Path(filename).stem)
    base = table_name
    suffix = 2
    while conn.execute("SELECT 1 FROM sqlite_master WHERE name=?", (table_name,)).fetchone():
        table_name = f"{base}_{suffix}"
        suffix += 1

    col_defs = []
    col_names: list[str] = []
    types = _infer_types(columns, data)
    for i, col in enumerate(columns):
        safe_col = sanitize_table_name(col) or f"col{i}"
        # Sanitising collapses distinct headers onto the same identifier
        # ("total sales" and "total-sales" both become "total_sales"), and
        # CREATE TABLE rejects a duplicate column outright. Suffix instead.
        if safe_col in col_names:
            n = 2
            while f"{safe_col}_{n}" in col_names:
                n += 1
            safe_col = f"{safe_col}_{n}"
        col_names.append(safe_col)
        col_defs.append(f'"{safe_col}" {types[i]}')

    conn.execute(f'CREATE TABLE "{table_name}" ({", ".join(col_defs)})')
    placeholders = ", ".join("?" for _ in col_names)
    quoted_cols = ", ".join('"%s"' % c for c in col_names)
    insert_sql = f'INSERT INTO "{table_name}" ({quoted_cols}) VALUES ({placeholders})'
    width = len(col_names)
    conn.executemany(
        insert_sql,
        ((list(row) + [None] * width)[:width] for row in data),
    )
    conn.commit()

    return {
        "table_name": table_name,
        "columns": col_names,
        "row_count": len(data),
        "source": Path(filename).name,
    }